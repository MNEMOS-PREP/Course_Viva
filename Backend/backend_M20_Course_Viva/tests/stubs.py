"""A model that answers the way the real one is asked to, by rule, so the
viva's tests run offline and say exactly what the model did."""

from __future__ import annotations

import io
import re
from typing import Any

from backend_M20_Course_Viva.llm import JOB_BANK, JOB_MARKING, LLMResult

#: Three pages of an Operating Systems lab, as slides (a slide is a page).
#: Each says when a page is evicted, so the stub's two points (evicted, memory
#: full) are on every page and its questions are long ones everywhere.
SLIDES = [
    "Page replacement decides which page to evict when memory is full and a new page is "
    "needed. FIFO replaces the page that has been in memory the longest. Belady's anomaly is a "
    "situation in which adding more frames increases the number of page faults. FIFO suffers "
    "from Belady's anomaly, while LRU does not.",
    "LRU replaces the page that has not been used for the longest time. It needs a counter or a "
    "stack to track how recently each page was used. The optimal algorithm replaces the page "
    "that will not be used for the longest time in the future; it cannot be built in practice "
    "and is used as a benchmark for the others. Every one of them acts only when memory is full "
    "and a page must be evicted.",
    "Aim: To simulate the FIFO and LRU page replacement algorithms and compare their page faults.\n"
    "Procedure: read the reference string, simulate each algorithm with three frames, and count "
    "a fault whenever the page is not in memory.\n"
    "Result: LRU produced fewer page faults than FIFO for the given reference string.\n"
    "Note: a page is evicted only when memory is full.",
]

#: An invented quote: on no page.
INVENTED = "Second chance gives every page a reference bit before it is evicted from memory"


def deck(slides: list[str] = SLIDES) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for text in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
        box.text_frame.text = text
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def _page_of(prompt: str) -> str:
    match = re.search(r"<<<PAGE\n(.*)\nPAGE$", prompt, re.S)
    return match.group(1) if match else ""


def _first_sentence(text: str) -> str:
    return re.split(r"(?<=[.!?])\s+", text.strip())[0]


class StubLLM:
    """`bank`: two questions per page — one quoting the page's first sentence,
    one quoting nothing on it (which the bank must drop). `marking`: a point
    is reached when the answer contains its `key` word; `wrong` answers name
    a wrong claim."""

    def __init__(self, *, available: bool = True, fail: set[str] | None = None):
        self._available = available
        self.fail = fail or set()
        self.calls: list[dict[str, Any]] = []

    @property
    def available(self) -> bool:
        return self._available

    def structured(self, *, job: str, system: str, user: str, schema: dict) -> LLMResult:
        self.calls.append({"job": job, "user": user})
        if job in self.fail:
            return LLMResult(ok=False, reason="scripted failure", job=job, model="stub")
        if job == JOB_BANK:
            page = _page_of(user)
            first = _first_sentence(page)
            return LLMResult(
                ok=True,
                job=job,
                model="stub-bank",
                data={
                    "questions": [
                        {
                            "family": "explain",
                            "question": f"Explain this: {' '.join(first.split()[:6])}",
                            "quote": first,
                            "answer_points": [
                                {"point": "the page is evicted", "probe": "Which page goes?"},
                                {"point": "memory is full", "probe": "When does it happen?"},
                            ],
                            "hint": "Think about what happens when there is no free frame.",
                            "difficulty": 1,
                        },
                        {
                            "family": "definition",
                            "question": "What is second chance replacement?",
                            "quote": INVENTED,
                            "answer_points": [{"point": "a reference bit", "probe": "What bit?"}],
                            "hint": "",
                            "difficulty": 2,
                        },
                    ]
                },
            )
        if job == JOB_MARKING:
            answer = re.search(r"<<<ANSWER\n(.*)\nANSWER$", user, re.S).group(1)
            judged = [int(i) for i in re.findall(r"^  (\d+)\. ", user, re.M)]
            points = []
            for index in judged:
                key = "evicted" if index == 0 else "full"
                hit = key in answer.lower()
                points.append({"index": index, "covered": hit, "evidence": answer if hit else ""})
            wrong = "it never evicts" in answer.lower()
            return LLMResult(
                ok=True,
                job=job,
                model="stub-marking",
                data={
                    "points": points,
                    "wrong_quote": "it never evicts" if wrong else "",
                    "wrong_why": "Pages are evicted when memory is full." if wrong else "",
                },
            )
        return LLMResult(ok=False, reason="unknown job", job=job)
