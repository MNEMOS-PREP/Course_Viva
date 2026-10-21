"""Text out of material, page by page (spec §6.3, FR-20.8).

A page keeps its number, because a viva question will cite it. What a format
calls a page: a PDF page, a slide, and for Word and plain text a ~3,000
character run (they have no pages of their own). A scan yields nothing and says
so (`no_text`, spec Trap 7) — the file is still stored and shown.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

#: Characters per "page" for formats without pages.
CHUNK = 3000
#: A document yielding less than this per page on average is read as a scan.
MIN_CHARS_PER_PAGE = 40


@dataclass
class Extracted:
    kind: str  # pdf | docx | pptx | text | image | unknown
    pages: list[str] = field(default_factory=list)
    status: str = "ok"  # ok | no_text | failed
    detail: str = ""


def sniff(data: bytes, filename: str) -> str:
    """The format from the bytes first, the name second: a renamed .docx is a zip."""
    name = filename.casefold()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if any(n.startswith("ppt/slides/") for n in names):
            return "pptx"
        return "unknown"
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image"
    if name.endswith((".txt", ".md", ".csv")):
        return "text"
    return "unknown"


MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "text": "text/plain",
    "image": "image/*",
    "unknown": "application/octet-stream",
}


def _chunks(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []
    pages, current = [], ""
    for para in text.split("\n"):
        if len(current) + len(para) > CHUNK and current:
            pages.append(current.strip())
            current = ""
        current += para + "\n"
    if current.strip():
        pages.append(current.strip())
    return pages


def extract(data: bytes, filename: str) -> Extracted:
    kind = sniff(data, filename)
    try:
        if kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "").strip() for page in reader.pages]
        elif kind == "docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            pages = _chunks("\n".join(parts))
        elif kind == "pptx":
            from pptx import Presentation

            deck = Presentation(io.BytesIO(data))
            pages = []
            for slide in deck.slides:
                texts = [
                    shape.text_frame.text
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
                ]
                pages.append("\n".join(texts).strip())
        elif kind == "text":
            pages = _chunks(data.decode("utf-8", "replace"))
        elif kind == "image":
            return Extracted(kind, [], "no_text", "an image: there is no OCR yet")
        else:
            return Extracted(kind, [], "failed", "not a PDF, Word, PowerPoint or text file")
    except Exception as exc:  # noqa: BLE001 - a broken file is reported, not raised
        return Extracted(kind, [], "failed", f"could not read it: {exc}")
    total = sum(len(p) for p in pages)
    if not pages or total < MIN_CHARS_PER_PAGE * max(len(pages), 1):
        return Extracted(kind, pages, "no_text", "little or no text: probably a scan")
    return Extracted(kind, pages, "ok")
